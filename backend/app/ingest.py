"""Document ingestion: parse -> clean -> chunk (with page/section metadata) -> embed -> index."""
import re
from pathlib import Path

import numpy as np

from . import retrieval
from .db import db

MAX_CHARS = 900
MIN_CHARS = 200


def clean(text: str) -> str:
    text = text.replace(" ", " ").replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80:
        return False
    if s.startswith("#"):
        return True
    if s.endswith((".", ",", ";")):
        return False
    words = s.split()
    if len(words) > 10:
        return False
    return s.isupper() or s.istitle() or s.endswith(":") or bool(re.match(r"^\d+(\.\d+)*[.)]?\s+\S", s))


def _heading_text(line: str) -> str:
    return re.sub(r"^#+\s*", "", line.strip()).rstrip(":").strip()


# ---------- parsers: return list of (page_or_None, text) ----------
def parse_pdf(path: Path):
    import pymupdf

    pages = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            pages.append((i, page.get_text("text")))
    return pages


def parse_docx(path: Path):
    import docx

    d = docx.Document(str(path))
    lines = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            lines.append("")
            continue
        if p.style is not None and p.style.name.lower().startswith(("heading", "title")):
            t = "# " + t
        lines.append(t)
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [(None, "\n".join(lines))]


def parse_text(path: Path):
    return [(None, path.read_text(encoding="utf-8", errors="ignore"))]


def parse_html(html: str):
    import trafilatura

    text = trafilatura.extract(html, include_tables=True, include_comments=False) or ""
    return [(None, text)]


def fetch_url(url: str):
    import httpx

    r = httpx.get(url, timeout=20, follow_redirects=True, headers={"User-Agent": "SupportAI/1.0"})
    r.raise_for_status()
    pages = parse_html(r.text)
    if not pages[0][1].strip():
        raise ValueError("No readable text found at that URL")
    return pages


# ---------- chunking ----------
def chunk_pages(pages):
    """Yield dicts {page, section, text}; section is the nearest preceding heading."""
    section = ""
    for page, raw in pages:
        text = clean(raw)
        if not text:
            continue
        out = []
        state = {"buf": [], "size": 0, "section": section}

        def flush():
            body = " ".join(state["buf"]).strip()
            if body:
                out.append({"page": page, "section": state["section"], "text": body})
            state["buf"], state["size"], state["section"] = [], 0, section

        def add(piece: str):
            if state["size"] and state["size"] + len(piece) > MAX_CHARS:
                flush()
            state["buf"].append(piece)
            state["size"] += len(piece) + 1

        for para in re.split(r"\n\s*\n", text):
            lines = [l for l in para.split("\n") if l.strip()]
            if not lines:
                continue
            if _is_heading(lines[0]):
                flush()
                section = _heading_text(lines[0])
                state["section"] = section
                lines = lines[1:]
                if not lines:
                    continue
            para_text = " ".join(l.strip() for l in lines)
            if len(para_text) > MAX_CHARS:
                for s in re.split(r"(?<=[.!?])\s+", para_text):
                    add(s)
            else:
                add(para_text)
        flush()

        merged = []
        for c in out:
            if (
                merged
                and len(c["text"]) < MIN_CHARS
                and merged[-1]["section"] == c["section"]
                and len(merged[-1]["text"]) + len(c["text"]) < MAX_CHARS * 1.3
            ):
                merged[-1]["text"] += " " + c["text"]
            else:
                merged.append(c)
        yield from merged


# ---------- persistence ----------
def index_document(tenant_id: int, name: str, source_type: str, source_ref: str, pages) -> int:
    chunks = list(chunk_pages(pages))
    with db() as conn:
        cur = conn.execute(
            "INSERT INTO documents(tenant_id,name,source_type,source_ref,status,chunk_count) VALUES(?,?,?,?,?,?)",
            (tenant_id, name, source_type, source_ref, "ready" if chunks else "failed", len(chunks)),
        )
        doc_id = cur.lastrowid
        if not chunks:
            conn.execute("UPDATE documents SET error=? WHERE id=?", ("No extractable text", doc_id))
            return doc_id
        embed_inputs = [f"{c['section']}. {c['text']}" if c["section"] else c["text"] for c in chunks]
        vecs = retrieval.embed_passages(embed_inputs)
        for c, v in zip(chunks, vecs):
            cc = conn.execute(
                "INSERT INTO chunks(tenant_id,doc_id,doc_name,page,section,text,embedding) VALUES(?,?,?,?,?,?,?)",
                (tenant_id, doc_id, name, c["page"], c["section"], c["text"], v.astype(np.float32).tobytes()),
            )
            fts_text = f"{c['section']} {c['text']}" if c["section"] else c["text"]
            conn.execute(
                "INSERT INTO chunks_fts(text,chunk_id,tenant_id) VALUES(?,?,?)",
                (fts_text, cc.lastrowid, tenant_id),
            )
    retrieval.invalidate(tenant_id)
    return doc_id


def ingest_file(tenant_id: int, path: Path, display_name: str) -> int:
    ext = path.suffix.lower()
    if ext == ".pdf":
        pages, kind = parse_pdf(path), "pdf"
    elif ext == ".docx":
        pages, kind = parse_docx(path), "docx"
    elif ext in (".txt", ".md", ".markdown", ".csv"):
        pages, kind = parse_text(path), "text"
    elif ext in (".html", ".htm"):
        pages, kind = parse_html(path.read_text(encoding="utf-8", errors="ignore")), "html"
    else:
        raise ValueError(f"Unsupported file type: {ext or 'unknown'}")
    return index_document(tenant_id, display_name, kind, display_name, pages)


def ingest_url(tenant_id: int, url: str) -> int:
    pages = fetch_url(url)
    return index_document(tenant_id, url, "url", url, pages)


def ingest_faq(tenant_id: int, question: str, answer: str) -> int:
    text = f"# {question.strip()}\n\n{answer.strip()}"
    return index_document(tenant_id, f"FAQ: {question.strip()[:60]}", "faq", question, [(None, text)])


def delete_document(tenant_id: int, doc_id: int) -> bool:
    with db() as conn:
        row = conn.execute("SELECT id FROM documents WHERE id=? AND tenant_id=?", (doc_id, tenant_id)).fetchone()
        if not row:
            return False
        for r in conn.execute("SELECT id FROM chunks WHERE doc_id=?", (doc_id,)).fetchall():
            conn.execute("DELETE FROM chunks_fts WHERE chunk_id=?", (r["id"],))
        conn.execute("DELETE FROM chunks WHERE doc_id=?", (doc_id,))
        conn.execute("DELETE FROM documents WHERE id=?", (doc_id,))
    retrieval.invalidate(tenant_id)
    return True
